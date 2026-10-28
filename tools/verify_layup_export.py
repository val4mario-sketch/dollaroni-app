"""Verifica che gli xlsx esportati diano gli stessi numeri con le formule UFFICIALI del template.

Per ogni file: sostituisce tutti i valori in cache delle formule con 12345, lo fa ricalcolare
a LibreOffice e confronta il risultato con i valori scritti dall'app. Se il motore dell'app
e le formule del Layup Calculator coincidono, lo scarto è ~1e-15.

Uso:  python tools/verify_layup_export.py out/layup-s1

Note (scoperte a fatica):
- LibreOffice di default NON ricalcola gli xlsx al caricamento: serve un profilo dedicato
  con OOXMLRecalcMode = 0 ("ricalcola sempre"). Lo creiamo in %TEMP%, senza toccare quello dell'utente.
- Con un percorso del profilo troppo lungo soffice si chiude con errore -1073740791: usiamo una cartella corta.
"""
import os
import pathlib
import re
import shutil
import subprocess
import sys
import zipfile

import openpyxl

SOFFICE = pathlib.Path(r"C:\Program Files\LibreOffice\program\soffice.com")
WORK = pathlib.Path(os.environ.get("TEMP", "/tmp")) / "dollaroni_lo"
SHEET = "Composite stacking calculation"
POISON = 12345
RECALC_ITEMS = (
    '<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>'
    '<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="ODFRecalcMode" oor:op="fuse"><value>0</value></prop></item>'
)


def soffice(*args):
    profile = "file:///" + str(WORK / "profile").replace("\\", "/")
    res = subprocess.run([str(SOFFICE), f"-env:UserInstallation={profile}", "--headless", *args], capture_output=True, text=True)
    if res.returncode != 0:
        sys.exit(f"soffice ha restituito {res.returncode}: {res.stderr or res.stdout}")


def ensure_profile(sample):
    reg = WORK / "profile" / "user" / "registrymodifications.xcu"
    if not reg.exists():
        soffice("--convert-to", "csv", "--outdir", str(WORK / "init"), str(sample))  # primo avvio: crea il profilo
    xml = reg.read_text(encoding="utf-8")
    if "OOXMLRecalcMode" not in xml:
        reg.write_text(xml.replace("</oor:items>", RECALC_ITEMS + "</oor:items>"), encoding="utf-8")


def poison_copy(src, dst):
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("xl/worksheets/"):
                data = re.sub(rb"(<f[^>]*>[^<]*</f>|<f[^>]*/>)<v>[^<]*</v>", rb"\1<v>%d</v>" % POISON, data)
            zout.writestr(item, data)


def main(folder):
    files = sorted(pathlib.Path(folder).glob("*.xlsx"))
    if not files:
        sys.exit(f"Nessun .xlsx in {folder}")
    for d in ("in", "out"):
        shutil.rmtree(WORK / d, ignore_errors=True)
        (WORK / d).mkdir(parents=True)
    for f in files:
        poison_copy(f, WORK / "in" / f.name)
    ensure_profile(WORK / "in" / files[0].name)
    soffice("--convert-to", "xlsx", "--outdir", str(WORK / "out"), *[str(WORK / "in" / f.name) for f in files])

    worst, failures = 0.0, 0
    for f in files:
        app = openpyxl.load_workbook(f, data_only=True)[SHEET]
        lo = openpyxl.load_workbook(WORK / "out" / f.name, data_only=True)[SHEET]
        for row in app.iter_rows():
            for c in row:
                a, b = c.value, lo[c.coordinate].value
                if b == POISON:
                    failures += 1
                    print(f"  {f.name} {c.coordinate}: non ricalcolata")
                elif isinstance(a, (int, float)) and isinstance(b, (int, float)):
                    worst = max(worst, abs(a - b) / max(1.0, abs(a)))
        print(f"{f.stem:45s} composito {lo['K27'].value:9.1f} g   fibre {lo['K28'].value:8.1f} g   resina {lo['K30'].value:8.1f} g")
    ok = failures == 0 and worst < 1e-9
    print(f"\nScarto massimo app vs formule ufficiali: {worst:.2e}  ->  {'OK' if ok else 'DIFFERENZE'}")
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "out/layup-s1")
